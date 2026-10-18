"""Scorer self-tests: the oracle (keys written as workbooks) scores perfectly, and
controlled damage to it costs exactly what the +1 / 0 / −3 rules say."""
import hashlib
from collections import Counter
from pathlib import Path

import openpyxl
import pytest
from openpyxl.comments import Comment

from carthagents.arbeitsversion import (OFFENE_PUNKTE, TEMPLATE, VERZEICHNIS, amount_columns, column_of,
                                        format_pointer, output_name, parse_pointer, write_workbook)
from carthagents.gt import load_doc
from carthagents.oracle import oracle_rows
from carthagents.scorer import score_doc

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "data" / "A3-Hackathon-Corpus"
needs_corpus = pytest.mark.skipif(not CORPUS.is_dir(), reason="sample corpus not present")


def _oracle(tmp_path, doc_id):
    doc = load_doc(next(CORPUS.rglob(f"{doc_id}.key.json")))
    for year in doc.years:
        write_workbook(tmp_path / output_name(doc_id, year), oracle_rows(doc, year), doc.period_end(year))
    return doc


def _tally(outcomes):
    return Counter((o.status, o.reason) for o in outcomes)


def _first_value_cell(ws, tab=VERZEICHNIS):
    col = amount_columns(tab)["ahk_open"]
    return ws[f"{col}2"]


def test_template_is_unchanged_copy():
    ref = ROOT / "data" / "Arbeitsversion-merged-layout-20.xlsx"
    if not ref.exists():
        pytest.skip("reference template not present")
    assert hashlib.sha256(TEMPLATE.read_bytes()).digest() == hashlib.sha256(ref.read_bytes()).digest()


def test_columns_come_from_feldkatalog():
    ws = openpyxl.load_workbook(TEMPLATE, read_only=True)[VERZEICHNIS]
    header = {c.column_letter: c.value for c in next(ws.iter_rows(max_row=1))}
    assert header[column_of(VERZEICHNIS, "ahk_01_01")] == "AHK 01.01."
    assert header[column_of(VERZEICHNIS, "bw_31_12_vj")] == "Buchwert 31.12. VJ"
    assert column_of("Gitter nach Anlagenklasse", "afa_gj") == "L"


@pytest.mark.parametrize("p", [{"sheet": "Kontenübersicht", "cell": "E9"},
                               {"page": 3, "bbox": [14.0, 81.8, 60.25, 90.1], "origin": "ocr"}])
def test_pointer_roundtrip(p):
    assert parse_pointer(format_pointer(p)) == p


@needs_corpus
@pytest.mark.parametrize("doc_id", ["f00_demo_01", "f00_demo_02", "f08_excel_spiegel_2y_01", "f03_block_spiegel_02"])
def test_oracle_is_perfect(tmp_path, doc_id):
    doc = _oracle(tmp_path, doc_id)
    out = score_doc(doc, tmp_path)
    assert out and all(o.status == "correct" for o in out)
    assert len(out) == sum(len(doc.expected(y)) for y in doc.years)


@needs_corpus
@pytest.mark.parametrize("doc_id", ["f00_demo_01", "f00_demo_02"])  # xlsx and pdf pointers
def test_damage_costs_what_the_rules_say(tmp_path, doc_id):
    doc = _oracle(tmp_path, doc_id)
    path = tmp_path / output_name(doc_id, doc.years[0])
    base = len(score_doc(doc, tmp_path))
    wb = openpyxl.load_workbook(path)
    ws = wb[VERZEICHNIS]
    cols = amount_columns(VERZEICHNIS)
    ws[f"{cols['ahk_open']}2"].value += 0.01                     # wrong value
    ws[f"{cols['zugang']}3"].comment = None                       # no pointer
    ws[f"{cols['abgang']}4"].value = None                         # left blank, unflagged
    ws[f"{cols['abgang']}5"].value = None                         # left blank, flagged
    wb[OFFENE_PUNKTE]["A4"], wb[OFFENE_PUNKTE]["B4"] = VERZEICHNIS, f"{cols['abgang']}5"
    ws[f"{cols['ahk_close']}6"].value = "12.345,67"               # text is not an amount
    moved = parse_pointer(ws[f"{cols['ahk_close']}8"].comment.text)  # pointer showing another row
    other = parse_pointer(ws[f"{cols['ahk_close']}20"].comment.text)
    assert moved != other
    ws[f"{cols['bw_close']}8"].comment = Comment(format_pointer(other), "t")
    ws[f"{cols['ahk_open']}400"] = 1.0                            # invented row
    ws[f"{cols['ahk_open']}400"].comment = Comment(format_pointer(other), "t")
    wb.save(path)
    t = _tally(score_doc(doc, tmp_path))
    assert t[("wrong", "value_mismatch")] == 2       # +0.01 and the text cell
    assert t[("wrong", "no_pointer")] == 1
    assert t[("blank", "missed")] == 1
    assert t[("blank", "flagged")] == 1
    assert t[("wrong", "pointer_mismatch")] == 1
    assert t[("wrong", "row_not_in_key")] == 1       # loses the vote to the real owner of row 20
    assert sum(t.values()) == base + 1               # only the invented value adds an outcome


@needs_corpus
def test_year_not_in_document_is_wrong(tmp_path):
    doc = _oracle(tmp_path, "f00_demo_01")
    src = tmp_path / output_name("f00_demo_01", doc.years[0])
    (tmp_path / output_name("f00_demo_01", 1999)).write_bytes(src.read_bytes())
    t = _tally(score_doc(doc, tmp_path))
    assert t[("wrong", "year_not_in_document")] == len(doc.expected(doc.years[0]))


@needs_corpus
def test_missing_workbook_is_blank(tmp_path):
    doc = load_doc(next(CORPUS.rglob("f00_demo_01.key.json")))
    out = score_doc(doc, tmp_path)
    assert out and {(o.status, o.reason) for o in out} == {("blank", "no_workbook")}
