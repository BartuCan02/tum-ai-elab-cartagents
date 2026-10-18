"""xlsx reader against the corpus answer keys (exact cell pointers and values)."""
from __future__ import annotations

import json
import re
from decimal import Decimal
from pathlib import Path

import pytest

from carthagents.grid import Cell, Document
from carthagents.readers.xlsx import format_number, read_xlsx

CORPUS = Path(__file__).resolve().parents[1] / "data" / "A3-Hackathon-Corpus"
XLSX = sorted(CORPUS.glob("f*/*.xlsx"))
_CACHE: dict[Path, Document] = {}


def _doc(path: Path) -> Document:
    if path not in _CACHE:
        _CACHE[path] = read_xlsx(path)
    return _CACHE[path]


def _index(doc: Document) -> dict[tuple[str, str], Cell]:
    return {(c.pointer["sheet"], c.pointer["cell"]): c for p in doc.pages for l in p.lines for c in l.cells}


def _parse_de(text: str) -> int | None:
    """Simple German amount parse to cents: '1.234,56-' / '(1.234,56)' / '-1.234,56' -> 123456 magnitude."""
    t = text.strip().strip("()").replace(" ", "").rstrip("-").lstrip("-+")
    if not re.fullmatch(r"\d{1,3}(\.\d{3})*(,\d+)?|\d+(,\d+)?", t):
        return None
    return int(Decimal(t.replace(".", "").replace(",", ".")) * 100)


@pytest.mark.parametrize("path", XLSX, ids=lambda p: p.stem)
def test_key_cells_reproduced(path: Path):
    key = json.loads((path.parent / f"{path.stem}.key.json").read_text())
    cells, bad, checked = _index(_doc(path)), [], 0
    for row in key["rows"]:
        ptr = row.get("pointer") or {}
        for field, coord in (ptr.get("cells") or {}).items():
            if field not in row.get("fields", {}):
                continue
            want, printed = abs(row["fields"][field]), (row.get("printed") or {}).get(field)
            if field in (row.get("flags") or {}):        # planted problem: key holds the truth, cell the print
                want = _parse_de(printed or "0")
            cell = cells.get((ptr["sheet"], coord))
            if cell is None:
                if printed not in ("", None) or want != 0:
                    bad.append(f"{ptr['sheet']}!{coord} {field}: no cell (printed {printed!r})")
                continue
            checked += 1
            # formula results carry float noise (2240.50000000003): compare at cent precision
            got = abs((cell.number * 100).quantize(1)) if cell.number is not None else _parse_de(cell.text)
            if got is None and cell.text in ("", "-") and want == 0:
                continue
            if got != want:
                bad.append(f"{ptr['sheet']}!{coord} {field}: got {cell.text!r}/{cell.number} want {want}")
    assert checked > 0
    assert not bad, "\n".join(bad[:20])


def test_hidden_column_flagged():
    doc = _doc(CORPUS / "f08_excel_spiegel_2y" / "f08_excel_spiegel_2y_07.xlsx")
    page = doc.pages[0]
    assert 7 in page.hidden_columns
    g = [c for l in page.lines for c in l.cells if c.pointer["cell"].startswith("G")]
    assert g and all(c.hidden for c in g)
    assert not any(c.hidden for l in page.lines for c in l.cells if c.pointer["cell"].startswith("F"))


@pytest.mark.parametrize("name", ["f09_excel_register_01", "f09_excel_register_05", "f09_excel_register_07"])
def test_hidden_rows_kept(name: str):
    doc = _doc(CORPUS / "f09_excel_register" / f"{name}.xlsx")
    hidden = [l for l in doc.pages[0].lines if l.hidden]
    assert len(hidden) >= 40
    assert all(c.hidden for l in hidden for c in l.cells)
    assert any(c.number is not None for l in hidden for c in l.cells)


def test_merged_header_span():
    doc = _doc(CORPUS / "f00_demo" / "f00_demo_01.xlsx")
    cells = _index(doc)
    head = cells[("Anlagenspiegel", "C5")]
    assert (head.x0, head.x1, head.y0, head.y1) == (2, 7, 4, 5)
    assert ("Anlagenspiegel", "D5") not in cells


def test_formulas_have_cached_numbers():
    doc = _doc(CORPUS / "f08_excel_spiegel_2y" / "f08_excel_spiegel_2y_01.xlsx")
    f = [c for p in doc.pages for l in p.lines for c in l.cells if c.formula]
    assert f and all(c.formula.startswith("=") for c in f)
    assert any(c.number is not None for c in f)


def test_numbers_as_text_stay_text():
    doc = _doc(CORPUS / "f09_excel_register" / "f09_excel_register_06.xlsx")
    cells = [c for p in doc.pages for l in p.lines for c in l.cells]
    assert all(c.number is None for c in cells)
    assert any(_parse_de(c.text) for c in cells)


def test_dates_rendered():
    doc = _doc(CORPUS / "f04_verzeichnis" / "f04_verzeichnis_03.xlsx")
    texts = [c.text for p in doc.pages for l in p.lines for c in l.cells]
    assert any(re.fullmatch(r"\d\d\.\d\d\.\d{4}", t) for t in texts)


@pytest.mark.parametrize("value,fmt,text", [
    ("1234.5", "#,##0.00", "1.234,50"),
    ("-1234.5", "#,##0.00;#,##0.00-", "1.234,50-"),
    ("-1234.5", "#,##0.00", "-1.234,50"),
    ("1234567", "#,##0", "1.234.567"),
    ("0", '_-* #,##0.00_-;\\-* #,##0.00_-;_-* \\-??_-;_-@_-', "-"),
    ("0", '#,##0.00;\\-#,##0.00;""', ""),
    ("0.125", "0.0%", "12,5%"),
    ("42", "General", "42"),
    ("3.25", "General", "3,25"),
])
def test_format_number(value: str, fmt: str, text: str):
    assert format_number(Decimal(value), fmt) == text


def test_never_raises_on_garbage(tmp_path: Path):
    bad = tmp_path / "x.xlsx"
    bad.write_bytes(b"not a zip")
    doc = read_xlsx(bad)
    assert doc.pages == [] and doc.notes


def test_synthetic_workbook(tmp_path: Path):
    """Formula without cached value, grouped hidden columns, numbers stored as text, hidden sheet."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"], ws["B1"], ws["C1"], ws["E1"] = "1.635,90", 1, 2, "=SUM(B1:C1)"   # openpyxl writes no cached value
    ws["B1"].number_format = "#,##0.00"
    ws.column_dimensions.group("B", "D", hidden=True)
    ws.merge_cells("A3:C4")
    ws["A3"] = "Kopf"
    wb.create_sheet("Versteckt").sheet_state = "hidden"
    wb["Versteckt"]["A1"] = True
    path = tmp_path / "s.xlsx"
    wb.save(path)
    doc = read_xlsx(path)
    cells = _index(doc)
    assert cells[("S", "A1")].number is None and cells[("S", "A1")].text == "1.635,90"
    assert cells[("S", "B1")].text == "1,00" and cells[("S", "B1")].number == 1
    f = cells[("S", "E1")]
    assert (f.text, f.number, f.formula) == ("", None, "=SUM(B1:C1)")
    assert doc.pages[0].hidden_columns >= {2, 3, 4} and cells[("S", "C1")].hidden and not f.hidden
    head = cells[("S", "A3")]
    assert (head.x0, head.x1, head.y0, head.y1) == (0, 3, 2, 4)
    assert [p.sheet for p in doc.pages] == ["S", "Versteckt"] and cells[("Versteckt", "A1")].text == "WAHR"
    assert any("formula without cached value" in n for n in doc.notes)
