"""Score filled Arbeitsversion workbooks against the answer keys (+1 / 0 / −3).

Only the Arbeitsversion is scored; the JSON audit trail is not read. Every amount
cell on the Anlagenverzeichnis and both Gitter tabs counts:

- correct (+1): the value equals the key and its pointer (cell comment) shows it
- blank (0):    a key value left empty (flagged if `Offene Punkte` names the cell)
- wrong (−3):   a different value, a missing/bad pointer, a value on the wrong tab,
                a value the document does not print, or a row matching no key row

Output rows are matched to key rows through their pointers, not through names or
numbers: a row belongs to the key row whose box (PDF) or cells (Excel) hold the
majority of its pointers. Output files: <doc_id>_<year>_Arbeitsversion.xlsx.

    python -m carthagents.scorer --keys data/A3-Hackathon-Corpus --pred outputs [--report reports/x]
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

from .arbeitsversion import (GITTER_FIRST, GITTER_TABS, OFFENE_PUNKTE, VERZEICHNIS, VERZEICHNIS_FIRST,
                             amount_columns, output_name, parse_pointer)
from .gt import Doc, KeyRow, load_corpus, pointer_in_row

POINTS = {"correct": 1, "blank": 0, "wrong": -3}


@dataclass
class Outcome:
    doc_id: str
    family: str
    format: str
    variant: str
    year: int
    tab: str
    cell: str
    field: str
    expected: int | None
    got: int | None
    status: str          # correct | blank | wrong
    reason: str          # detail, e.g. value_mismatch, pointer_mismatch, missed, flagged


@dataclass
class PredValue:
    cell: str
    cents: int | None    # None: the cell holds something that is not an exact amount
    pointer: dict | None


def _cents(v) -> int | None:
    if isinstance(v, bool):
        return None
    try:
        d = (Decimal(v) if isinstance(v, int) else Decimal(repr(v)) if isinstance(v, float) else None)
    except InvalidOperation:
        return None
    if d is None:
        return None
    d *= 100
    return int(d) if d == d.to_integral_value() else None


def read_prediction(path: Path) -> tuple[dict[tuple[str, int], dict[str, PredValue]], set[tuple[str, str]]]:
    """(tab, row) -> field -> value, plus the (tab, cell) pairs listed under Offene Punkte."""
    wb = openpyxl.load_workbook(path)
    rows: dict[tuple[str, int], dict[str, PredValue]] = defaultdict(dict)
    for tab in (VERZEICHNIS, *GITTER_TABS):
        ws = wb[tab]
        first = VERZEICHNIS_FIRST if tab == VERZEICHNIS else GITTER_FIRST
        for f, col in amount_columns(tab).items():
            for (cell,) in ws[f"{col}{first}:{col}{max(ws.max_row, first)}"]:
                v = cell.value
                if v is None or (isinstance(v, str) and (v.startswith("=") or not v.strip())):
                    continue  # empty or a template formula: nothing was written
                rows[(tab, cell.row)][f] = PredValue(cell.coordinate, _cents(v),
                                                     parse_pointer(cell.comment.text if cell.comment else None))
    flagged = set()
    for a, b, *_ in wb[OFFENE_PUNKTE].iter_rows(min_row=4, max_col=4, values_only=True):
        if a and b:
            flagged.add((str(a).strip(), str(b).strip()))
    wb.close()
    return rows, flagged


def _match(pred: dict[tuple[str, int], dict[str, PredValue]], keyrows: list[KeyRow]) -> dict[tuple[str, int], KeyRow]:
    best: dict[tuple[str, int], tuple[KeyRow, int]] = {}
    for prow, values in pred.items():
        votes = Counter()
        for v in values.values():
            for kr in keyrows:
                if pointer_in_row(v.pointer, kr):
                    votes[kr.index] += 1
        top = votes.most_common(2)
        if top and (len(top) == 1 or top[0][1] > top[1][1]):
            best[prow] = (next(k for k in keyrows if k.index == top[0][0]), top[0][1])
    # one output row per key row: the strongest claim wins, ties go to the first row
    owner: dict[int, tuple[str, int]] = {}
    for prow, (kr, n) in sorted(best.items(), key=lambda x: (-x[1][1], x[0][1])):
        owner.setdefault(kr.index, prow)
    return {prow: kr for prow, (kr, _) in best.items() if owner[kr.index] == prow}


def score_doc(doc: Doc, pred_dir: Path) -> list[Outcome]:
    m = doc.meta
    out: list[Outcome] = []

    def emit(year, tab, cell, field, expected, got, status, reason):
        out.append(Outcome(doc.doc_id, m["family"], m["format"], m["variant"], year, tab, cell, field,
                           expected, got, status, reason))

    for path in sorted(pred_dir.glob(f"{doc.doc_id}_*_Arbeitsversion.xlsx")):
        y = re.fullmatch(re.escape(doc.doc_id) + r"_(\d{4})_Arbeitsversion\.xlsx", path.name)
        if y and int(y.group(1)) not in doc.years:
            pred, _ = read_prediction(path)
            for (tab, _row), values in pred.items():
                for f, v in values.items():
                    emit(int(y.group(1)), tab, v.cell, f, None, v.cents, "wrong", "year_not_in_document")
    for year in doc.years:
        expected = doc.expected(year)
        path = pred_dir / output_name(doc.doc_id, year)
        if not path.exists():
            for e in expected:
                emit(year, e.row.tab, "", e.field, e.cents, None, "blank", "no_workbook")
            continue
        pred, flagged = read_prediction(path)
        keyrows = [r for r in doc.rows if r.year == year]
        matched = _match(pred, keyrows)
        seen: set[int] = set()
        for prow, values in pred.items():
            tab, rownum = prow
            kr = matched.get(prow)
            if kr is None:
                reason = "no_pointer" if all(v.pointer is None for v in values.values()) else "row_not_in_key"
                for f, v in values.items():
                    emit(year, tab, v.cell, f, None, v.cents, "wrong", reason)
                continue
            if kr.tab != tab:
                reason = "row_not_scored_on_any_tab" if kr.tab is None else f"belongs_on_{kr.tab}"
                for f, v in values.items():
                    emit(year, tab, v.cell, f, kr.fields.get(f), v.cents, "wrong", reason)
                continue
            seen.add(kr.index)
            cols = amount_columns(tab)
            for e in (e for e in expected if e.row.index == kr.index):
                v = values.pop(e.field, None)
                cell = f"{cols[e.field]}{rownum}"
                if v is None:
                    emit(year, tab, cell, e.field, e.cents, None, "blank",
                         "flagged" if (tab, cell) in flagged else "missed")
                elif not pointer_in_row(v.pointer, kr, e.field):
                    emit(year, tab, v.cell, e.field, e.cents, v.cents, "wrong",
                         "no_pointer" if v.pointer is None else "pointer_mismatch")
                elif v.cents != e.cents:
                    emit(year, tab, v.cell, e.field, e.cents, v.cents, "wrong", "value_mismatch")
                else:
                    emit(year, tab, v.cell, e.field, e.cents, v.cents, "correct", "")
            for f, v in values.items():
                emit(year, tab, v.cell, f, None, v.cents, "wrong", "field_not_printed")
        for e in expected:
            if e.row.index not in seen:
                emit(year, e.row.tab, "", e.field, e.cents, None, "blank", "row_missed")
    return out


# --- reporting -------------------------------------------------------------------------
def summarize(outcomes: list[Outcome], by: str | None = None) -> dict[str, dict]:
    groups: dict[str, Counter] = defaultdict(Counter)
    for o in outcomes:
        for g in ("ALL", getattr(o, by) if by else None):
            if g is not None:
                groups[str(g)][o.status] += 1
    table = {}
    for g, c in groups.items():
        n = c["correct"] + c["blank"] + c["wrong"]
        scored = c["correct"] + c["wrong"]
        table[g] = {"values": n, "correct": c["correct"], "blank": c["blank"], "wrong": c["wrong"],
                    "score": sum(POINTS[s] * c[s] for s in POINTS),
                    "precision": round(c["correct"] / scored, 4) if scored else None,
                    "coverage": round(c["correct"] / n, 4) if n else None}
    return table


def format_table(table: dict[str, dict], title: str) -> str:
    head = f"{title:<26}{'values':>8}{'correct':>9}{'blank':>8}{'wrong':>8}{'score':>9}{'prec':>8}{'cover':>8}"
    lines = [head, "-" * len(head)]
    for g in sorted(table, key=lambda g: (g == "ALL", g)):
        r = table[g]
        pct = lambda x: "-" if x is None else f"{100 * x:.1f}%"
        lines.append(f"{g:<26}{r['values']:>8}{r['correct']:>9}{r['blank']:>8}{r['wrong']:>8}{r['score']:>9}"
                     f"{pct(r['precision']):>8}{pct(r['coverage']):>8}")
    return "\n".join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Score Arbeitsversion outputs against answer keys")
    ap.add_argument("--keys", required=True, type=Path, help="folder with *.key.json (searched recursively)")
    ap.add_argument("--pred", required=True, type=Path, help="folder with <doc_id>_<year>_Arbeitsversion.xlsx")
    ap.add_argument("--report", type=Path, help="write summary.json and outcomes.csv here")
    ap.add_argument("--only", nargs="*", help="score only these doc_ids or families")
    args = ap.parse_args(argv)
    docs = [d for d in load_corpus(args.keys)
            if not args.only or d.doc_id in args.only or d.meta["family"] in args.only]
    if not docs:
        print("no answer keys found", file=sys.stderr)
        return 1
    outcomes = [o for d in docs for o in score_doc(d, args.pred)]
    excluded = sum((d.excluded for d in docs), Counter())
    for by in ("family", "format", "variant"):
        print(format_table(summarize(outcomes, by), by), end="\n\n")
    reasons = Counter((o.status, o.reason) for o in outcomes if o.status != "correct")
    print("not correct, by reason:")
    for (s, r), n in reasons.most_common():
        print(f"  {s:<6} {r:<32}{n:>8}")
    print(f"\nkey values not scored (no home in the Arbeitsversion): {sum(excluded.values())}")
    for r, n in excluded.most_common():
        print(f"  {r:<60}{n:>8}")
    if args.report:
        args.report.mkdir(parents=True, exist_ok=True)
        summary = {by: summarize(outcomes, by) for by in ("family", "format", "variant", "tab", "field")}
        summary["reasons"] = {f"{s}:{r}": n for (s, r), n in reasons.items()}
        summary["excluded"] = dict(excluded)
        (args.report / "summary.json").write_text(json.dumps(summary, indent=1, ensure_ascii=False), encoding="utf-8")
        with open(args.report / "outcomes.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(Outcome.__dataclass_fields__))
            w.writeheader()
            w.writerows(asdict(o) for o in outcomes if o.status != "correct")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
