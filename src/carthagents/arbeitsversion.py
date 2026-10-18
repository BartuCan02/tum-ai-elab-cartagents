"""The official Arbeitsversion template: field mapping, pointer notation and writer.

The template (resources/Arbeitsversion.xlsx) is ground truth and is never modified;
every output is a filled copy of it, one workbook per document and fiscal year.
Column letters are resolved through the template's own Feldkatalog sheets, so the
code only knows field IDs, never hard-coded positions.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from functools import lru_cache
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

TEMPLATE = Path(__file__).with_name("resources") / "Arbeitsversion.xlsx"

VERZEICHNIS = "Anlagenverzeichnis"
GITTER_KLASSE = "Gitter nach Anlagenklasse"
GITTER_POSITION = "Gitter nach Bilanzposition"
OFFENE_PUNKTE = "Offene Punkte"
GITTER_TABS = (GITTER_KLASSE, GITTER_POSITION)
CATALOG = {VERZEICHNIS: "Feldkatalog Anlagenverzeichnis", GITTER_KLASSE: "Feldkatalog Gitter",
           GITTER_POSITION: "Feldkatalog Gitter"}

# Canonical amount fields (answer-key names) -> Feld-ID per tab. A field missing from a
# tab's map has no column there and is kept in the JSON only (e.g. the asset-level
# depreciation total, whose template columns are split into planmäßig/außerplanmäßig).
_COMMON = {
    "ahk_open": "ahk_01_01", "zugang": "ahk_zugang", "abgang": "ahk_abgang", "umbuchung": "ahk_umbuchung",
    "ahk_close": "ahk_31_12", "kum_afa_open": "afa_01_01", "afa_abgang": "afa_abgang",
    "zuschreibung": "zuschreibung", "kum_afa_close": "afa_31_12", "bw_close": "bw_31_12_gj",
    "bw_prev": "bw_31_12_vj",
}
FIELD_IDS = {
    VERZEICHNIS: dict(_COMMON),
    GITTER_KLASSE: {**_COMMON, "afa_gj": "afa_gj", "bew_ahk": "ahk_bewegung_netto", "bew_wb": "afa_bewegung_netto"},
    GITTER_POSITION: {**_COMMON, "afa_gj": "afa_gj", "bew_ahk": "ahk_bewegung_netto", "bew_wb": "afa_bewegung_netto"},
}
VERZEICHNIS_FIRST, VERZEICHNIS_LAST = 2, 501  # template: "keine Eintragung unterhalb von Zeile 501"
GITTER_FIRST, GITTER_SUMME = 8, 28           # data rows 8–27, row 28 "Summe laut Mandant"


def output_name(stem: str, year: int) -> str:
    return f"{stem}_{year}_Arbeitsversion.xlsx"


@lru_cache(maxsize=None)
def _catalog(template: str) -> dict[str, dict[str, str]]:
    wb = openpyxl.load_workbook(template, read_only=True)
    out = {}
    for sheet in set(CATALOG.values()):
        out[sheet] = {str(r[1]).strip(): str(r[0]).strip() for r in wb[sheet].iter_rows(min_row=5, max_col=2, values_only=True)
                      if r[0] and r[1]}
    wb.close()
    return out


def column_of(tab: str, feld_id: str, template: Path = TEMPLATE) -> str:
    """Column letter of a Feld-ID on a tab, read from the template's Feldkatalog."""
    return _catalog(str(template))[CATALOG[tab]][feld_id]


def amount_columns(tab: str, template: Path = TEMPLATE) -> dict[str, str]:
    """Canonical field -> column letter for every amount column of a tab."""
    return {f: column_of(tab, fid, template) for f, fid in FIELD_IDS[tab].items()}


# --- pointer notation in cell comments -------------------------------------------------
# PDF:   "Herkunft: Seite 2, Box 14.00 81.80 60.20 90.10 (text_layer)"
# Excel: "Herkunft: 'Tabelle1'!C8"
_PDF_RE = re.compile(r"Herkunft: Seite (\d+), Box ([-\d.]+) ([-\d.]+) ([-\d.]+) ([-\d.]+)(?: \((\w+)\))?")
_XLSX_RE = re.compile(r"Herkunft: '(.*)'!([A-Z]+\d+)")


def format_pointer(p: dict) -> str:
    if p.get("cell"):
        return f"Herkunft: '{p['sheet']}'!{p['cell']}"
    x0, y0, x1, y1 = p["bbox"]
    return f"Herkunft: Seite {p['page']}, Box {x0:.2f} {y0:.2f} {x1:.2f} {y1:.2f} ({p.get('origin', 'text_layer')})"


def parse_pointer(text: str | None) -> dict | None:
    if not text:
        return None
    if m := _XLSX_RE.search(text):
        return {"sheet": m.group(1), "cell": m.group(2)}
    if m := _PDF_RE.search(text):
        return {"page": int(m.group(1)), "bbox": [float(m.group(i)) for i in range(2, 6)], "origin": m.group(6)}
    return None


# --- writer ----------------------------------------------------------------------------
@dataclass
class Value:
    cents: int | None          # None = left blank on purpose; then `flag` says why
    pointer: dict | None = None
    flag: str | None = None


@dataclass
class OutRow:
    tab: str
    values: dict[str, Value]                         # canonical field -> value
    text: dict[str, object] = field(default_factory=dict)  # Feld-ID -> printed text (Stammdaten)
    total: bool = False                              # a client total ("Summe …"), written at the Summe row


@dataclass
class OpenPoint:
    tab: str
    where: str
    question: str
    status: str = "offen"


def write_workbook(path: Path, rows: list[OutRow], period_end: date | None, open_points: list[OpenPoint] = (),
                   template: Path = TEMPLATE) -> None:
    wb = openpyxl.load_workbook(template)
    flags = list(open_points)
    for tab in GITTER_TABS:
        ws = wb[tab]
        if period_end:
            ws["C2"] = period_end
        data = [r for r in rows if r.tab == tab and not r.total]
        totals = [r for r in rows if r.tab == tab and r.total]
        extra = max(0, len(data) - (GITTER_SUMME - GITTER_FIRST))
        if extra:
            ws.insert_rows(GITTER_SUMME, extra)
        summe = GITTER_SUMME + extra
        if len(totals) > 1:
            ws.insert_rows(summe + 1, len(totals) - 1)
        for i, r in enumerate(data):
            flags += _fill(ws, GITTER_FIRST + i, r, tab, template)
        for i, r in enumerate(totals):
            flags += _fill(ws, summe + i, r, tab, template)
    ws = wb[VERZEICHNIS]
    assets = [r for r in rows if r.tab == VERZEICHNIS]
    if len(assets) > VERZEICHNIS_LAST - VERZEICHNIS_FIRST + 1:
        raise ValueError(f"{len(assets)} assets exceed the template's {VERZEICHNIS_LAST - 1} rows")
    for i, r in enumerate(assets):
        flags += _fill(ws, VERZEICHNIS_FIRST + i, r, VERZEICHNIS, template)
    if flags:
        ws = wb[OFFENE_PUNKTE]
        ws["A4"] = None
        for i, p in enumerate(flags):
            for col, v in zip("ABCD", (p.tab, p.where, p.question, p.status)):
                ws[f"{col}{4 + i}"] = v
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _fill(ws, row: int, r: OutRow, tab: str, template: Path) -> list[OpenPoint]:
    flags = []
    for fid, text in r.text.items():
        ws[f"{column_of(tab, fid, template)}{row}"] = text
    cols = amount_columns(tab, template)
    for f, v in r.values.items():
        if f not in cols:
            continue
        cell = ws[f"{cols[f]}{row}"]
        if v.cents is None:
            if v.flag:
                flags.append(OpenPoint(tab, cell.coordinate, v.flag))
            continue
        cell.value = v.cents / 100  # exact: shortest float repr round-trips the cents
        if v.pointer:
            cell.comment = Comment(format_pointer(v.pointer), "carthagents")
    return flags
