"""xlsx -> grid.Document. Reports cells as Excel shows them; never computes formulas.

Every worksheet becomes one Page (including hidden sheets; the interpreter decides relevance).
Numbers carry their exact cached value (`Cell.number`), the formula text (`Cell.formula`) and the
displayed text rendered from the number format in German locale (`1.234,56`). Strings that merely
look like numbers stay text. Hidden rows/columns are kept and flagged, never dropped.
"""
from __future__ import annotations

import datetime as dt
import re
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from ..grid import Cell, Document, Line, Page


# ---------------------------------------------------------------- number formats

_SECTION_SPLIT = re.compile(r'((?:"[^"]*"|\\.|[^;])*)(?:;|$)')


def _sections(fmt: str) -> list[str]:
    parts = [m.group(1) for m in _SECTION_SPLIT.finditer(fmt)]
    while len(parts) > 1 and parts[-1] == "" and not fmt.endswith(";"):
        parts.pop()
    return parts[:4] or [""]


def _strip(section: str) -> tuple[str, str, str]:
    """Split a format section into (prefix literal, numeric pattern, suffix literal)."""
    section = re.sub(r"\[[^\]]*\]", "", section)            # [Red], [$-407], conditions
    out, pat, i, seen_digit = ["", ""], [], 0, False
    while i < len(section):
        ch = section[i]
        tgt = 1 if seen_digit else 0
        if ch == '"':
            j = section.find('"', i + 1)
            j = len(section) if j < 0 else j
            out[tgt] += section[i + 1:j]
            i = j + 1
            continue
        if ch == "\\" and i + 1 < len(section):
            out[tgt] += section[i + 1]
            i += 2
            continue
        if ch in "_*":                                        # padding / fill: invisible
            i += 2
            continue
        if ch in "0#?,." and (ch not in ",." or pat):
            pat.append(ch)
            seen_digit = True
        elif ch == "%":
            out[tgt] += "%"
        elif ch in "-+ ()/$€:":
            out[tgt] += ch
        i += 1
    return out[0].strip(), "".join(pat), out[1].strip()


def _group(digits: str) -> str:
    rev = digits[::-1]
    return ".".join(rev[i:i + 3] for i in range(0, len(rev), 3))[::-1]


def _general(value: Decimal) -> str:
    if value == value.to_integral_value():
        return str(int(value))
    text = format(value.normalize(), "f")
    if len(text.replace("-", "").replace(".", "")) > 11:     # Excel General shows ~10 significant digits
        text = format(value.quantize(Decimal(1).scaleb(-max(0, 10 - len(str(int(abs(value)))))),
                                     ROUND_HALF_UP).normalize(), "f")
    return text.replace(".", ",")


def format_number(value: Decimal, fmt: str | None) -> str:
    """Render a number as Excel would display it with format `fmt` in German locale."""
    fmt = fmt or "General"
    if fmt.lower() == "general":
        return _general(value)
    secs = _sections(fmt)
    if value < 0 and len(secs) >= 2:
        sec, val, sign = secs[1], -value, ""
    elif value == 0 and len(secs) >= 3:
        sec, val, sign = secs[2], value, ""
    else:
        sec, val, sign = secs[0], value, "-" if value < 0 else ""
        val = abs(val)
    if sec.lower() == "general":
        return sign + _general(val)
    pre, pat, post = _strip(sec)
    if "@" in sec and not pat:
        return (pre + post).strip()
    if not pat:                                               # literal-only section ("" or "-")
        return (pre + post).strip()
    if "%" in pre + post:
        val *= 100
    int_pat, _, dec_pat = pat.partition(".")
    while int_pat.endswith(","):                              # trailing comma scales by 1000
        int_pat, val = int_pat[:-1], val / 1000
    places = sum(1 for c in dec_pat if c in "0#?")
    q = val.quantize(Decimal(1).scaleb(-places), ROUND_HALF_UP)
    int_digits, _, frac = format(q, "f").partition(".")
    if int_digits == "0" and "0" not in int_pat:
        int_digits = ""
    if "," in int_pat:
        int_digits = _group(int_digits)
    min_frac = sum(1 for c in dec_pat if c == "0")
    frac = frac.rstrip("0").ljust(min_frac, "0") if places else ""
    body = int_digits + ("," + frac if frac else "")
    if q == 0:
        sign = ""
    return (sign + pre + body + post).strip()


def _format_date(value: dt.date | dt.time) -> str:
    if isinstance(value, dt.datetime) and value.time() != dt.time(0) and value.date() == dt.date(1899, 12, 30):
        return value.strftime("%H:%M:%S")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    return value.strftime("%d.%m.%Y")


# ---------------------------------------------------------------- hidden rows / columns

def _hidden_columns(ws) -> frozenset[int]:
    cols: set[int] = set()
    for key, dim in ws.column_dimensions.items():
        if not dim.hidden:
            continue
        lo = dim.min or column_index_from_string(key)
        cols.update(range(lo, max(lo, dim.max or lo) + 1))
    return frozenset(cols)


# ---------------------------------------------------------------- reader

def _cell(fc, vc, sheet: str, coord: str, box: tuple[int, int, int, int], hidden: bool,
          notes: list[str]) -> Cell | None:
    """Build one grid cell from the formula-view cell `fc` and the cached-value cell `vc`."""
    raw, val = fc.value, (vc.value if vc is not None else None)
    formula = None
    if fc.data_type == "f" or (isinstance(raw, str) and raw.startswith("=")):
        formula = raw if isinstance(raw, str) else getattr(raw, "text", str(raw))
    elif raw is not None and not isinstance(raw, (str, int, float, bool, dt.date, dt.time, Decimal)):
        formula = str(getattr(raw, "text", raw))              # array formula / data table
    if formula is None:
        val = raw
    if val is None and formula is None:
        return None
    number, text = None, ""
    try:
        if isinstance(val, bool):
            text = "WAHR" if val else "FALSCH"
        elif isinstance(val, (int, float, Decimal)):
            number = Decimal(val) if isinstance(val, (int, Decimal)) else Decimal(repr(val))
            if not number.is_finite():
                number, text = None, str(val)
            else:
                text = format_number(number, fc.number_format)
        elif isinstance(val, (dt.date, dt.time)):
            text = _format_date(val)
        elif val is not None:
            text = str(val).strip()
    except Exception as exc:                                  # never raise on odd formats
        notes.append(f"{sheet}!{coord}: could not render value {val!r} ({exc})")
        text = str(val).strip()
    if not text and number is None and formula is None:
        return None
    if formula is not None and val is None:
        notes.append(f"{sheet}!{coord}: formula without cached value: {formula}")
    c0, r0, c1, r1 = box
    return Cell(text=text, pointer={"sheet": sheet, "cell": coord}, x0=c0 - 1, x1=c1, y0=r0 - 1, y1=r1,
                number=number, formula=formula, hidden=hidden)


def _read_sheet(ws, wv, index: int, notes: list[str]) -> Page:
    hidden_cols = _hidden_columns(ws)
    hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
    spans: dict[tuple[int, int], tuple[int, int]] = {}        # top-left (row, col) -> (max_row, max_col)
    covered: set[tuple[int, int]] = set()
    for rng in ws.merged_cells.ranges:
        spans[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
        covered.update((r, c) for r in range(rng.min_row, rng.max_row + 1)
                       for c in range(rng.min_col, rng.max_col + 1) if (r, c) != (rng.min_row, rng.min_col))
    rows: dict[int, list[Cell]] = {}
    for row in ws.iter_rows():
        for fc in row:
            r, c = getattr(fc, "row", None), getattr(fc, "column", None)
            if r is None or (r, c) in covered or fc.value is None:
                continue
            coord = f"{get_column_letter(c)}{r}"
            r1, c1 = spans.get((r, c), (r, c))
            hidden = r in hidden_rows or c in hidden_cols
            vc = wv[coord] if wv is not None else None
            cell = _cell(fc, vc, ws.title, coord, (c, r, c1, r1), hidden, notes)
            if cell is not None:
                rows.setdefault(r, []).append(cell)
    lines = [Line(cells=sorted(cells, key=lambda x: x.x0), y0=r - 1, y1=r, row=r, hidden=r in hidden_rows)
             for r, cells in sorted(rows.items())]
    if ws.sheet_state != "visible":
        notes.append(f"sheet {ws.title!r} is {ws.sheet_state}")
    return Page(index=index, lines=lines, origin="excel_cell", sheet=ws.title,
                width=float(ws.max_column or 0), height=float(ws.max_row or 0), hidden_columns=hidden_cols)


def read_xlsx(path: Path) -> Document:
    """Read every worksheet of an xlsx file into grid pages. Problems go to Document.notes."""
    doc = Document(path=str(path), format="xlsx", pages=[])
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as exc:
        doc.notes.append(f"cannot open workbook: {exc}")
        return doc
    try:
        wbv = load_workbook(path, data_only=True)
    except Exception as exc:
        wbv = None
        doc.notes.append(f"cannot read cached values: {exc}")
    for i, ws in enumerate(wb.worksheets, start=1):
        try:
            wv = wbv[ws.title] if wbv is not None and ws.title in wbv.sheetnames else None
            doc.pages.append(_read_sheet(ws, wv, i, doc.notes))
        except Exception as exc:
            doc.notes.append(f"sheet {ws.title!r}: read failed: {exc}")
    for ws in wb.chartsheets if hasattr(wb, "chartsheets") else []:
        doc.notes.append(f"chartsheet {ws.title!r} skipped")
    return doc
